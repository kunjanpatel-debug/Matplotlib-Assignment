
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

slide = prs.slides.add_slide(blank_layout)

title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.8))
tf = title_box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Session 11: Matplotlib High-Resolution Chart Export Presentation"
p.font.name = "Calibri"
p.font.size = Pt(22)
p.font.bold = True
p.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

# 
sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(11.7), Inches(0.5))
tf2 = sub_box.text_frame
p2 = tf2.paragraphs[0]
p2.text = "Demonstrating Programmatic Image Embedding: Zomato Orders (PNG @ 150 DPI) & Weekly Steps Line Plot"
p2.font.name = "Calibri"
p2.font.size = Pt(13)
p2.font.italic = True
p2.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

img_path1 = 'session11_task1_zomato_orders.png'
if os.path.exists(img_path1):
    slide.shapes.add_picture(img_path1, Inches(0.8), Inches(2.0), width=Inches(5.7))

img_path2 = 'session11_task2_step_count.png'
if os.path.exists(img_path2):
    slide.shapes.add_picture(img_path2, Inches(6.8), Inches(2.0), width=Inches(5.7))

prs.save("session11_task4_presentation.pptx")
print("session11_task4_embed_ppt.py executed successfully. Saved presentation as session11_task4_presentation.pptx.")
